from __future__ import annotations

import csv
from dataclasses import dataclass
from pathlib import Path
from typing import Callable, Optional

import cv2
import numpy as np
from PIL import Image
from pptx import Presentation
from pptx.util import Inches


@dataclass
class ConversionConfig:
    sample_fps: float = 3.0
    stable_seconds: float = 1.0
    change_threshold: float = 0.035
    similarity: float = 0.94
    min_gap: float = 1.5
    max_width: int = 1920
    crop: tuple = (0, 0, 1, 1)


def resize_frame(frame, max_width):
    h, w = frame.shape[:2]
    if w <= max_width:
        return frame
    scale = max_width / w
    return cv2.resize(frame, (max_width, int(h * scale)), interpolation=cv2.INTER_AREA)


def crop_frame(frame, crop):
    h, w = frame.shape[:2]
    l, t, r, b = crop
    x1, y1 = int(l * w), int(t * h)
    x2, y2 = int(r * w), int(b * h)
    if x2 <= x1 or y2 <= y1:
        raise ValueError("잘못된 crop 설정입니다.")
    return frame[y1:y2, x1:x2]


def normalized_gray(frame):
    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    small = cv2.resize(gray, (320, 180), interpolation=cv2.INTER_AREA)
    return cv2.GaussianBlur(small, (5, 5), 0)


def diff_score(a, b):
    return float(np.mean(cv2.absdiff(a, b)) / 255.0)


def phash(frame):
    gray = cv2.cvtColor(frame, cv2.COLOR_BGR2GRAY)
    small = cv2.resize(gray, (32, 32), interpolation=cv2.INTER_AREA).astype(np.float32)
    dct = cv2.dct(small)
    low = dct[:8, :8].flatten()
    median = np.median(low[1:])
    bits = low > median
    value = 0
    for bit in bits:
        value = (value << 1) | int(bit)
    return value


def phash_similarity(a, b):
    return 1.0 - ((a ^ b).bit_count() / 64.0)


def save_pdf(image_paths, pdf_path):
    if not image_paths:
        return
    images = []
    for p in image_paths:
        with Image.open(p) as im:
            images.append(im.convert("RGB").copy())
    images[0].save(pdf_path, save_all=True, append_images=images[1:], resolution=150)
    for im in images:
        im.close()


def save_pptx(image_paths, pptx_path):
    prs = Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank = prs.slide_layouts[6]
    slide_ratio = prs.slide_width / prs.slide_height

    for p in image_paths:
        slide = prs.slides.add_slide(blank)
        with Image.open(p) as im:
            iw, ih = im.size
        ratio = iw / ih
        if ratio > slide_ratio:
            width = prs.slide_width
            height = int(width / ratio)
            left = 0
            top = int((prs.slide_height - height) / 2)
        else:
            height = prs.slide_height
            width = int(height * ratio)
            left = int((prs.slide_width - width) / 2)
            top = 0
        slide.shapes.add_picture(str(p), left, top, width=width, height=height)

    prs.save(pptx_path)


class VideoToSlides:
    def __init__(self, config: ConversionConfig):
        self.cfg = config

    def run(self, video: Path, output_dir: Path,
            progress: Optional[Callable[[float, str], None]] = None):
        video = Path(video)
        output_dir = Path(output_dir)
        image_dir = output_dir / "images"
        image_dir.mkdir(parents=True, exist_ok=True)

        def report(value, msg):
            if progress:
                progress(float(value), msg)

        cap = cv2.VideoCapture(str(video))
        if not cap.isOpened():
            raise RuntimeError(
                "영상을 열 수 없습니다. MP4 코덱 또는 FFmpeg 설치 상태를 확인하세요."
            )

        fps = cap.get(cv2.CAP_PROP_FPS) or 30.0
        total = int(cap.get(cv2.CAP_PROP_FRAME_COUNT) or 0)
        duration = total / fps if total else 0
        step = max(1, round(fps / self.cfg.sample_fps))
        stable_needed = max(1, round(self.cfg.stable_seconds * self.cfg.sample_fps))

        previous = None
        stable_count = 0
        last_saved = -1e9
        candidates = []
        idx = 0

        report(1, "영상 분석을 시작합니다...")

        while True:
            ok, frame = cap.read()
            if not ok:
                break

            idx += 1
            if idx % step:
                continue

            timestamp = idx / fps
            frame = crop_frame(frame, self.cfg.crop)
            small = normalized_gray(frame)

            if previous is None:
                previous = small
                stable_count = 1
                candidates.append((timestamp, frame.copy()))
                continue

            score = diff_score(previous, small)

            if score >= self.cfg.change_threshold:
                stable_count = 0
            else:
                stable_count += 1

            if stable_count >= stable_needed and timestamp - last_saved >= self.cfg.min_gap:
                candidates.append((timestamp, frame.copy()))
                last_saved = timestamp
                stable_count = 0

            previous = small

            if duration:
                report(min(65, 5 + timestamp / duration * 60),
                       f"분석 중... {timestamp/60:.1f}분 / {duration/60:.1f}분")

        cap.release()

        report(68, f"후보 프레임 {len(candidates)}개를 찾았습니다. 중복 제거 중...")

        accepted = []
        accepted_hashes = []

        for timestamp, frame in candidates:
            h = phash(frame)
            if any(phash_similarity(h, old) >= self.cfg.similarity for old in accepted_hashes):
                continue

            frame = resize_frame(frame, self.cfg.max_width)
            number = len(accepted) + 1
            path = image_dir / f"slide_{number:04d}.png"
            rgb = cv2.cvtColor(frame, cv2.COLOR_BGR2RGB)
            Image.fromarray(rgb).save(str(path), format="PNG")

            accepted.append((timestamp, path))
            accepted_hashes.append(h)

        if not accepted:
            raise RuntimeError(
                "슬라이드를 찾지 못했습니다. 분석 FPS를 높이거나 변화 임계값을 낮춰보세요."
            )

        report(80, f"{len(accepted)}개 슬라이드 확정. PPTX 생성 중...")
        pptx = output_dir / f"{video.stem}_slides.pptx"
        save_pptx([p for _, p in accepted], pptx)

        report(91, "PDF 생성 중...")
        pdf = output_dir / f"{video.stem}_slides.pdf"
        save_pdf([p for _, p in accepted], pdf)

        csv_path = output_dir / f"{video.stem}_slides.csv"
        with csv_path.open("w", newline="", encoding="utf-8-sig") as f:
            writer = csv.writer(f)
            writer.writerow(["slide_number", "timestamp_seconds", "image_path"])
            for i, (timestamp, path) in enumerate(accepted, 1):
                writer.writerow([i, f"{timestamp:.3f}", str(path)])

        report(100, "모든 작업이 완료되었습니다.")

        return {
            "count": len(accepted),
            "pptx": str(pptx),
            "pdf": str(pdf),
            "images": str(image_dir),
            "csv": str(csv_path),
        }
